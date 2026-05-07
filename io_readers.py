from __future__ import annotations
import os
import re
import unicodedata
import pandas as pd

class InputError(Exception):
    pass


def detect_header_by_required_columns(
    df_preview: pd.DataFrame,
    required_columns: list[str],
    max_scan_rows: int = 50
) -> int:
    """
    Detecta la fila de encabezados buscando coincidencia
    con columnas esperadas (Oracle-style reports).
    """
    required_norm = {c.lower().strip() for c in required_columns}

    for idx in range(min(len(df_preview), max_scan_rows)):
        row_vals = (
            df_preview.iloc[idx]
            .dropna()
            .astype(str)
            .str.lower()
            .str.strip()
            .tolist()
        )

        matches = sum(1 for v in row_vals if v in required_norm)

        # si detecta varias columnas clave → es el header correcto
        if matches >= max(2, len(required_norm) // 2):
            return idx

    return -1


def read_table(
    path: str,
    skiprows: int | None = None,
    sheet_name: int | str = 0,
    required_columns_hint: list[str] | None = None
) -> pd.DataFrame:

    if not os.path.exists(path):
        raise InputError(f"No se encontró el archivo: {path}")

    ext = os.path.splitext(path)[1].lower()

    try:
        if ext in (".xlsx", ".xls"):
            df_preview = pd.read_excel(
                path,
                sheet_name=sheet_name,
                header=None,
                nrows=60
            )

            if skiprows is not None:
                header_row = skiprows

            elif required_columns_hint:
                header_row = detect_header_by_required_columns(
                    df_preview,
                    required_columns_hint
                )
                if header_row < 0:
                    raise InputError(
                        f"No se pudo detectar encabezado usando columnas esperadas en {path}"
                    )
            else:
                raise InputError(
                    f"Debe indicarse skiprows o required_columns_hint para {path}"
                )

            df = pd.read_excel(
                path,
                sheet_name=sheet_name,
                header=header_row
            )

        elif ext == ".csv":
            df = pd.read_csv(path)

        else:
            raise InputError(f"Formato no soportado: {ext}")

    except Exception as e:
        raise InputError(f"Error leyendo {path}: {e}")

    # limpieza
    df = df.loc[:, ~df.columns.isna()]
    df = df.dropna(axis=1, how="all")

    if df.empty:
        raise InputError(f"{path} quedó vacío tras lectura.")

    return df

def _fold(s: str) -> str:
    s = "" if s is None else str(s)
    s = s.strip().lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = re.sub(r"\s+", " ", s)
    return s

def _parse_number_loose(x):
    if x is None:
        return None
    # openpyxl puede devolver float/int ya
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip()
    if not s:
        return None
    s = s.replace("\u00A0", " ").replace(" ", "")
    # patrón miles/decimales
    if "." in s and "," in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    else:
        if "," in s and "." not in s:
            s = s.replace(",", ".")
        elif "." in s and s.count(".") > 1:
            s = s.replace(".", "")
    s = re.sub(r"[^0-9\.\-]", "", s)
    try:
        return float(s)
    except Exception:
        return None

def extract_labeled_amount_excel(
    path: str,
    label: str,
    sheet_name: int | str = 0,
    search_max_cols: int = 60
) -> tuple[float | None, str | None]:
    """
    Busca 'label' en cualquier celda de la hoja y devuelve:
      (valor_numérico, dirección_celda_label)
    tomando el primer número a la derecha en la misma fila.
    """
    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise InputError(f"openpyxl no disponible para extraer etiquetas: {e}")

    if not os.path.exists(path):
        raise InputError(f"No se encontró el archivo: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if isinstance(sheet_name, str) else wb.worksheets[int(sheet_name)]

    target = _fold(label)

    for row in ws.iter_rows():
        # limitamos columnas para no recorrer 16k columnas si existieran
        cells = row[:search_max_cols]
        for j, cell in enumerate(cells):
            v = cell.value
            if v is None:
                continue
            if target in _fold(v):
                # buscar a la derecha un número
                for k in range(j + 1, min(len(cells), j + 15)):
                    num = _parse_number_loose(cells[k].value)
                    if num is not None:
                        addr = cell.coordinate
                        return num, addr

    return None, None
