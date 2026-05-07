from __future__ import annotations

from datetime import datetime
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.worksheet import Worksheet


def _autosize(ws: Worksheet, max_width: int = 60):
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        lengths = []
        for cell in col_cells:
            if cell.value is None:
                continue
            lengths.append(len(str(cell.value)))
        if not lengths:
            continue
        width = min(max(lengths) + 2, max_width)
        ws.column_dimensions[col_letter].width = width


def _format_sheet(ws: Worksheet):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # encabezado en negrita
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    _autosize(ws)


def write_report(
    out_path: str,
    *,
    readme_rows: list[tuple[str, str]],
    resumen_df: pd.DataFrame,
    sheets: dict[str, pd.DataFrame]
):
    """
    Escribe Excel con:
      - README
      - Resumen
      - hojas por regla
    """
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # README
        df_readme = pd.DataFrame(readme_rows, columns=["Parámetro", "Valor"])
        df_readme.to_excel(writer, sheet_name="README", index=False)

        # Resumen
        resumen_df.to_excel(writer, sheet_name="Resumen", index=False)

        # Hojas de hallazgos
        for name, df in sheets.items():
            if df is None:
                df = pd.DataFrame({"Sin datos": []})
            if df.empty:
                # hoja con mensaje para que siempre exista
                df2 = pd.DataFrame({"Mensaje": ["Sin hallazgos para esta regla."]})
                df2.to_excel(writer, sheet_name=name[:31], index=False)
            else:
                df.to_excel(writer, sheet_name=name[:31], index=False)

        # Aplicar formato al final
        wb = writer.book
        for ws in wb.worksheets:
            _format_sheet(ws)

        # Un toque de estilo en README
        ws_readme = wb["README"]
        ws_readme["A1"].font = Font(bold=True)
        ws_readme["B1"].font = Font(bold=True)
        _autosize(ws_readme)

def generar_reporte_inconsistencias(df_validado):

    df = df_validado.copy()

    inconsistencias = df[
        (~df["Debito OK"]) |
        (~df["Credito OK"]) |
        (~df["Saldo OK"])
    ]

    return inconsistencias.sort_values("Cuenta Contable")